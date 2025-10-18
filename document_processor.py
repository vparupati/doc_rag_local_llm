"""
Document processing module for handling various file formats.
"""
import os
import logging
from pathlib import Path
from typing import List, Dict, Any
import PyPDF2
import docx
from langchain.text_splitter import RecursiveCharacterTextSplitter
from langchain.schema import Document

# Configure logging
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

class DocumentProcessor:
    """Handles document processing and text extraction."""
    
    def __init__(self, chunk_size: int = 1000, chunk_overlap: int = 200):
        self.chunk_size = chunk_size
        self.chunk_overlap = chunk_overlap
        self.text_splitter = RecursiveCharacterTextSplitter(
            chunk_size=chunk_size,
            chunk_overlap=chunk_overlap,
            length_function=len,
        )
    
    def extract_text_from_pdf(self, file_path: str) -> str:
        """Extract text from PDF file."""
        try:
            text = ""
            with open(file_path, 'rb') as file:
                pdf_reader = PyPDF2.PdfReader(file)
                for page in pdf_reader.pages:
                    text += page.extract_text() + "\n"
            logger.info(f"Successfully extracted text from PDF: {file_path}")
            return text
        except Exception as e:
            logger.error(f"Error extracting text from PDF {file_path}: {str(e)}")
            raise
    
    def extract_text_from_txt(self, file_path: str) -> str:
        """Extract text from plain text file."""
        try:
            with open(file_path, 'r', encoding='utf-8') as file:
                text = file.read()
            logger.info(f"Successfully extracted text from TXT: {file_path}")
            return text
        except Exception as e:
            logger.error(f"Error extracting text from TXT {file_path}: {str(e)}")
            raise
    
    def extract_text_from_docx(self, file_path: str) -> str:
        """Extract text from DOCX file."""
        try:
            doc = docx.Document(file_path)
            text = ""
            for paragraph in doc.paragraphs:
                text += paragraph.text + "\n"
            logger.info(f"Successfully extracted text from DOCX: {file_path}")
            return text
        except Exception as e:
            logger.error(f"Error extracting text from DOCX {file_path}: {str(e)}")
            raise
    
    def extract_text_from_md(self, file_path: str) -> str:
        """Extract text from Markdown file."""
        try:
            with open(file_path, 'r', encoding='utf-8') as file:
                text = file.read()
            logger.info(f"Successfully extracted text from MD: {file_path}")
            return text
        except Exception as e:
            logger.error(f"Error extracting text from MD {file_path}: {str(e)}")
            raise
    
    def extract_text(self, file_path: str) -> str:
        """Extract text from file based on extension."""
        file_extension = Path(file_path).suffix.lower()
        
        if file_extension == '.pdf':
            return self.extract_text_from_pdf(file_path)
        elif file_extension == '.txt':
            return self.extract_text_from_txt(file_path)
        elif file_extension == '.docx':
            return self.extract_text_from_docx(file_path)
        elif file_extension == '.md':
            return self.extract_text_from_md(file_path)
        else:
            raise ValueError(f"Unsupported file format: {file_extension}")
    
    def process_document(self, file_path: str, metadata: Dict[str, Any] = None) -> List[Document]:
        """Process a document and return chunked documents."""
        try:
            # Extract text
            text = self.extract_text(file_path)
            
            # Create base metadata
            base_metadata = {
                "source": file_path,
                "filename": Path(file_path).name,
                "file_type": Path(file_path).suffix.lower(),
            }
            
            # Add custom metadata if provided
            if metadata:
                base_metadata.update(metadata)
            
            # Split text into chunks
            texts = self.text_splitter.split_text(text)
            
            # Create Document objects
            documents = []
            for i, text_chunk in enumerate(texts):
                doc_metadata = base_metadata.copy()
                doc_metadata["chunk_id"] = i
                doc_metadata["chunk_size"] = len(text_chunk)
                
                documents.append(Document(
                    page_content=text_chunk,
                    metadata=doc_metadata
                ))
            
            logger.info(f"Processed document {file_path} into {len(documents)} chunks")
            return documents
            
        except Exception as e:
            logger.error(f"Error processing document {file_path}: {str(e)}")
            raise
    
    def process_multiple_documents(self, file_paths: List[str]) -> List[Document]:
        """Process multiple documents and return all chunked documents."""
        all_documents = []
        
        for file_path in file_paths:
            try:
                documents = self.process_document(file_path)
                all_documents.extend(documents)
            except Exception as e:
                logger.error(f"Failed to process {file_path}: {str(e)}")
                continue
        
        logger.info(f"Processed {len(file_paths)} files into {len(all_documents)} total chunks")
        return all_documents
